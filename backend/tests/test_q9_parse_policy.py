"""
q9 deterministic parsing contract tests (owner directive v1.0).

Covers:
  - the unsupported_by_policy terminal contract shape
  - policy invariants (no OCR, no parser AI models)
  - pdf_requires_ocr_rejection gate (scanned/mixed PDF vs digital text PDF,
    non-PDF passthrough)
  - worker terminal marking (_mark_ingest_unsupported_by_policy)
  - terminal-stage exclusion sets across pipeline/summary/readiness/control-plane

Unit tests mock the Mongo boundary; no live stores are contacted.
"""

from __future__ import annotations

from dataclasses import dataclass
from unittest.mock import AsyncMock, MagicMock

import pytest

from models.schemas import SourceTier
from services.ingestion import parse_policy


def _fake_pdf_result(text: str, pages: list[str] | None = None):
    """Minimal parse-result stand-in for the adapter's usability heuristic."""
    from services.ingestion.docling_adapter import DoclingParseResult, Section

    return DoclingParseResult(
        text=text,
        markdown=text,
        sections=[],
        pages=pages if pages is not None else ([text] if text else []),
        has_structure=False,
        source_tier=SourceTier.ocr_ast,
        num_pages=max(1, len(pages or [text])),
        source_format="pypdf_fast_text",
        filename="scan.pdf",
    )


# ── Contract shape ──────────────────────────────────────────────────────────


def test_unsupported_contract_matches_owner_spec():
    contract = parse_policy.unsupported_by_policy_result()
    assert contract["status"] == "unsupported_by_policy"
    assert contract["reason"] == "document_requires_ocr"
    assert contract["retryable"] is False
    assert contract["ocr_enabled"] is False
    assert contract["ocr_supported"] is False
    assert contract["parser_model_loads"] == 0


def test_policy_invariants_are_absolute():
    assert parse_policy.OCR_SUPPORTED is False
    assert parse_policy.OCR_ENABLED is False
    assert parse_policy.OCR_RUNTIME_DEPENDENCIES == 0
    assert parse_policy.VISION_MODELS_LOADED_BY_PARSER == 0
    assert parse_policy.PARSER_AI_MODELS == 0


def test_supported_formats_are_deterministic_parsers():
    # Every supported format maps to a deterministic parser; no OCR/vision
    # entry may ever appear.
    assert set(parse_policy.SUPPORTED_FORMATS) == {
        "markdown",
        "txt",
        "html",
        "pdf",
        "docx",
        "epub",
        "json",
        "yaml",
    }
    for parser in parse_policy.SUPPORTED_FORMATS.values():
        assert "ocr" not in parser.lower()
        assert "vision" not in parser.lower()


# ── Rejection gate ──────────────────────────────────────────────────────────


def test_scanned_pdf_is_rejected():
    result = _fake_pdf_result("", pages=["", "", ""])
    assert parse_policy.pdf_requires_ocr_rejection(result, "scan.pdf", "application/pdf")


def test_sparse_mixed_pdf_is_rejected():
    # A handful of stray characters across many pages — a mixed PDF whose
    # required content cannot be extracted without OCR.
    pages = ["p", "", "42", "", "", ""]
    result = _fake_pdf_result(" ".join(pages), pages=pages)
    assert parse_policy.pdf_requires_ocr_rejection(
        result, "mixed.pdf", "application/pdf"
    )


def test_digital_text_pdf_is_not_rejected():
    body = "Benesh movement notation records choreography on a five-line stave. "
    pages = [body * 20, body * 20, body * 20]
    result = _fake_pdf_result("".join(pages), pages=pages)
    assert not parse_policy.pdf_requires_ocr_rejection(
        result, "book.pdf", "application/pdf"
    )


def test_non_pdf_never_hits_the_ocr_gate():
    result = _fake_pdf_result("", pages=[])
    assert not parse_policy.pdf_requires_ocr_rejection(
        result, "notes.md", "text/markdown"
    )
    assert not parse_policy.pdf_requires_ocr_rejection(result, "book.epub", "")


# ── Worker terminal marking ─────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_mark_ingest_unsupported_by_policy_sets_terminal_stage():
    db = MagicMock()
    db.__getitem__ = MagicMock(
        return_value=MagicMock(update_one=AsyncMock(return_value=None))
    )

    reason = await __import__(
        "services.ingestion.worker", fromlist=["_mark_ingest_unsupported_by_policy"]
    )._mark_ingest_unsupported_by_policy(
        db=db, doc_id="doc1", corpus_id="corp1"
    )

    assert "OCR" in reason
    coll = db["documents"]
    coll.update_one.assert_awaited_once()
    query, update = coll.update_one.await_args.args
    assert query == {"doc_id": "doc1", "corpus_id": "corp1"}
    set_clause = update["$set"]
    assert set_clause["ingest_stage"] == "unsupported_by_policy"
    assert set_clause["queryable"] is False
    assert set_clause["excluded_from_readiness"] is True
    assert set_clause["unsupported_policy"]["status"] == "unsupported_by_policy"
    assert set_clause["unsupported_policy"]["retryable"] is False
    assert set_clause["unsupported_policy"]["ocr_enabled"] is False


# ── Terminal exclusion sets ─────────────────────────────────────────────────


def test_terminal_sets_cover_unsupported_by_policy():
    from services.ingestion import document_pipeline_jobs, readiness, summary_jobs
    from services.control_plane import desired_state

    for stage_set in (
        document_pipeline_jobs.TERMINAL_SKIP_INGEST_STAGES,
        summary_jobs.TERMINAL_SKIP_INGEST_STAGES,
        readiness.EXCLUDED_DOCUMENT_STAGES,
        desired_state.EXCLUDED_DOCUMENT_STAGES,
    ):
        assert "unsupported_by_policy" in stage_set
        # Existing terminal exclusions must be preserved.
        assert "skipped_duplicate" in stage_set
        assert "skipped_nonsemantic" in stage_set


# ── pdf-inspector integration ───────────────────────────────────────────────


def _minimal_text_pdf_bytes() -> bytes:
    """Hand-rolled one-page digital text PDF (no external fixtures)."""
    # Body must clear the deterministic usability thresholds
    # (>=1200 total chars) so a real text PDF is never mistaken for a scan.
    sentence = (
        "Benesh movement notation records choreography on a five-line stave "
        "for stage works and preserves limb positions through symbolic marks. "
    )
    body = sentence * 12
    content = (
        b"BT /F1 12 Tf 50 700 Td (" + body.encode("utf-8") + b") Tj ET"
    )
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 6\n0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(xref).encode() + b"\n%%EOF"
    return out


@pytest.mark.asyncio
async def test_text_pdf_parses_through_pdf_inspector():
    """The pinned deterministic engine owns text-PDF parsing end to end."""
    from services.ingestion.docling_adapter import parse_document

    result = await parse_document(
        _minimal_text_pdf_bytes(),
        filename="benesh.pdf",
        mime="application/pdf",
        do_ocr=False,
    )
    assert result.source_format == "pdf_inspector"
    assert "Benesh" in (result.markdown or "")
    # Deterministic contract: the parser never loads models or runs OCR.
    assert parse_policy.PARSER_AI_MODELS == 0
    assert parse_policy.OCR_ENABLED is False


def test_text_pdf_fails_the_usability_gate_only_when_sparse():
    """The same engine output feeds the parse-policy usability gate."""
    from services.ingestion.docling_adapter import parse_document
    import asyncio

    result = asyncio.get_event_loop().run_until_complete(
        parse_document(
            _minimal_text_pdf_bytes(),
            filename="benesh.pdf",
            mime="application/pdf",
        )
    )
    assert not parse_policy.pdf_requires_ocr_rejection(
        result, "benesh.pdf", "application/pdf"
    )


# ── q9-6: native EPUB + DOCX deterministic adapters ────────────────────────


def test_epub_and_docx_map_to_native_deterministic_parsers():
    # Owner directive: EPUB = deterministic zip/XHTML, DOCX = native OOXML.
    # pdf-inspector is scoped to PDFs only and must not appear here.
    assert parse_policy.SUPPORTED_FORMATS["epub"] == "deterministic_zip_xhtml_parser"
    assert parse_policy.SUPPORTED_FORMATS["docx"] == "deterministic_native_ooxml_parser"
    for fmt in ("epub", "docx"):
        assert "inspector" not in parse_policy.SUPPORTED_FORMATS[fmt]
        assert "pdf" not in parse_policy.SUPPORTED_FORMATS[fmt]


@pytest.mark.asyncio
async def test_epub_native_parse_never_invokes_pdf_inspector(monkeypatch):
    """q9 pdf_parser_scope: EPUB parses natively; pdf-inspector is PDF-only."""
    pytest.importorskip("ebooklib")
    from ebooklib import epub
    from services.ingestion import docling_adapter as adapter

    def _forbidden(*_a, **_k):
        raise AssertionError("pdf-inspector must not run for EPUB")

    monkeypatch.setattr(adapter, "_parse_pdf_with_inspector", _forbidden)

    book = epub.EpubBook()
    book.set_identifier("q9-epub")
    book.set_title("Movement Notation Primer")
    book.set_language("en")
    book.add_author("Q. Author")
    chapter = epub.EpubHtml(title="Stave", file_name="stave.xhtml", lang="en")
    # Body must clear the adapter's >=200-char usability guard so a real
    # book is never mistaken for a navigation-only EPUB.
    chapter.content = (
        "<h1>Stave Basics</h1>"
        "<p>The five-line stave anchors every Benesh symbol and records the "
        "dancer's limb positions across time for faithful stage reconstruction.</p>"
    ) * 6
    book.add_item(chapter)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.toc = (chapter,)
    book.spine = ["nav", chapter]
    import io

    buffer = io.BytesIO()
    epub.write_epub(buffer, book)

    result = await adapter.parse_document(
        raw_bytes=buffer.getvalue(),
        filename="primer.epub",
        mime="application/epub+zip",
        do_ocr=False,
    )
    assert result.source_format == "local_epub"
    assert result.title == "Movement Notation Primer"
    assert "stave" in (result.text or "").lower()


@pytest.mark.asyncio
async def test_docx_native_parse_never_invokes_pdf_inspector(monkeypatch):
    """q9 pdf_parser_scope: DOCX parses via OOXML; pdf-inspector is PDF-only."""
    pytest.importorskip("docx")
    from docx import Document
    from services.ingestion import docling_adapter as adapter

    def _forbidden(*_a, **_k):
        raise AssertionError("pdf-inspector must not run for DOCX")

    monkeypatch.setattr(adapter, "_parse_pdf_with_inspector", _forbidden)

    doc = Document()
    doc.add_heading("Notation Overview", level=1)
    doc.add_paragraph("Benesh notation captures movement on a stave.")
    doc.add_heading("Symbols", level=2)
    doc.add_paragraph("Limb positions are encoded as discrete marks.")
    import io

    buffer = io.BytesIO()
    doc.save(buffer)

    result = await adapter.parse_document(
        raw_bytes=buffer.getvalue(),
        filename="overview.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        do_ocr=False,
    )
    assert result.source_format == "local_docx"
    headings = {
        s.text.strip() for s in result.sections if s.element_type == "section_heading"
    }
    assert "Notation Overview" in headings
    assert "Symbols" in headings
